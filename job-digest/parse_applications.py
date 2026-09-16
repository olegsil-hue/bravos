#!/usr/bin/env python3
"""Parse hh.ru responses export (xlsx) into application records."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path


UI_NOISE = {
    "Только непрочитанные",
    "Бета-версия",
    "Поиск с ИИ-помощником",
}

TODAY = date(2026, 9, 16)


def excel_date(value, today: date = TODAY) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 40000 < value < 50000:
        return (datetime(1899, 12, 30) + timedelta(days=int(value))).date()
    if not isinstance(value, str):
        return None
    text = value.strip().lower().replace("\xa0", " ")
    if text == "сегодня":
        return today
    if text == "вчера":
        return today - timedelta(days=1)
    if text == "позавчера":
        return today - timedelta(days=2)
    match = re.fullmatch(r"(\d+)\s+дн(?:ень|я|ей)\s+назад", text)
    if match:
        return today - timedelta(days=int(match.group(1)))
    match = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", text)
    if match:
        day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
        if year < 100:
            year += 2000
        return date(year, month, day)
    return None


def is_date(value) -> bool:
    return excel_date(value) is not None


def classify_status(raw: str) -> str:
    text = (raw or "").replace("\xa0", " ").strip()
    low = text.lower()
    if text in {"Отказ"} or "закрыли эту позицию" in low or "компания заблокирована" in low:
        return "Отказ"
    if text in {"Отклик на вакансию", "Отклик"}:
        return "Отклик"
    if text in {"Приглашение", "Собеседование", "Входящий звонок"} or "собеседован" in low:
        return "Собеседование"
    if text in {"Не просмотрено", "Просмотрено"}:
        return "Отклик"
    if len(text) > 40 or text.startswith(("Олег", "Добрый", "Здравствуйте", "Привет", "Спасибо")):
        if any(word in low for word in ("отказ", "не рассматриваем", "не подходит", "закрыли")):
            return "Отказ"
        return "Ответ работодателя"
    if text.startswith("Директор ") or text.startswith("Руководитель "):
        # Next vacancy title leaked into the status cell.
        return "Отклик"
    return "Отклик" if not text else text


def parse_xlsx(path: str | Path, today: date = TODAY) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    values = [row[0] for row in workbook.active.iter_rows(values_only=True) if row[0] is not None]
    if values and values[0] == "Только непрочитанные":
        values = values[1:]

    records: list[dict] = []
    used: set[int] = set()
    index = 1
    while index < len(values) - 1:
        if not is_date(values[index]) or (index - 1) in used:
            index += 1
            continue
        title = values[index - 1]
        if not isinstance(title, str) or title in UI_NOISE or is_date(title):
            index += 1
            continue
        company = values[index + 1] if index + 1 < len(values) else None
        status_raw = values[index + 2] if index + 2 < len(values) else ""
        if not isinstance(company, str) or is_date(company) or company in UI_NOISE:
            index += 1
            continue
        parsed_date = excel_date(values[index], today)
        status_text = status_raw.replace("\xa0", " ").strip() if isinstance(status_raw, str) else ""
        records.append(
            {
                "title": title.replace("\xa0", " ").strip(),
                "date": parsed_date.isoformat() if parsed_date else str(values[index]),
                "company": company.replace("\xa0", " ").strip(),
                "status": classify_status(status_text),
                "status_raw": status_text[:300],
            }
        )
        used.update({index - 1, index, index + 1})
        if index + 2 < len(values):
            used.add(index + 2)
        index += 3
    return records


def company_key(name: str) -> str:
    text = (name or "").lower().replace("ё", "е")
    text = re.sub(r"[\"«»“”]", "", text)
    text = re.sub(
        r"\b(ооо|ао|пао|зао|оао|ип|llc|ltd|inc|группа компаний|гк|компания)\b",
        " ",
        text,
    )
    return re.sub(r"[^a-zа-я0-9]+", "", text)


def company_summary(records: list[dict]) -> list[dict]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        grouped[record["company"]].append(record)

    rank = {"Собеседование": 4, "Ответ работодателя": 3, "Отклик": 2, "Отказ": 1}

    rows = []
    for company, items in grouped.items():
        items = sorted(items, key=lambda item: item["date"], reverse=True)
        latest = items[0]
        statuses = {item["status"] for item in items}
        had_rejection = "Отказ" in statuses
        latest_non_reject = next((item for item in items if item["status"] != "Отказ"), None)
        if had_rejection and (latest["status"] == "Отказ" or latest_non_reject is None):
            bucket = "отказ"
            display_status = "Отказ"
        elif had_rejection:
            bucket = "отказ"
            display_status = f"{latest['status']} (ранее отказ)"
        else:
            bucket = "отклик"
            display_status = max(items, key=lambda item: rank.get(item["status"], 0))["status"]
        rows.append(
            {
                "company": company,
                "key": company_key(company),
                "date": latest["date"],
                "status": display_status,
                "bucket": bucket,
                "count": len(items),
                "last_title": latest["title"],
            }
        )
    rows.sort(key=lambda row: (row["date"], row["company"]), reverse=True)
    return rows


def main() -> None:
    source = Path("/home/ubuntu/.cursor/projects/workspace/uploads/______________01.26____16.09_bf2b.xlsx")
    here = Path(__file__).resolve().parent
    records = parse_xlsx(source)
    companies = company_summary(records)
    (here / "applications.json").write_text(
        json.dumps({"records": records, "companies": companies}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"records={len(records)} companies={len(companies)}")
    print("buckets", {bucket: sum(1 for row in companies if row["bucket"] == bucket) for bucket in ("отклик", "отказ")})


if __name__ == "__main__":
    main()
